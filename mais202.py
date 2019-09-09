
#--- GATHER DATA FROM CSV FILES ---#

import csv 
import array as arr
from collections import defaultdict

ownershipData = defaultdict(list) # each value in each column is appended to a list

with open('home_ownership_data.csv') as f:
    reader = csv.DictReader(f) # read rows into a dictionary format
    for row in reader: 
        for (a,b) in row.items(): # go over each column name and value 
            ownershipData[a].append(b)
                                 

loanData = defaultdict(list) # each value in each column is appended to a list

with open('loan_data.csv') as f:
    reader = csv.DictReader(f) # read rows into a dictionary format
    for row in reader: 
        for (a,b) in row.items(): # go over each column name and value 
            loanData[a].append(b)
                                 

#values to be sumed up by type
mortgage = 0
rent = 0
own = 0

#number of the ownership type in list
numMortgage = 0
numRent = 0
numOwn = 0

for i in ownershipData['member_id']:
	
	if ownershipData['home_ownership'][ownershipData['member_id'].index(i)] == 'MORTGAGE':
		mortgage += int(loanData['loan_amnt'][loanData['member_id'].index(i)])
		numMortgage+=1
		
	elif ownershipData['home_ownership'][ownershipData['member_id'].index(i)] == 'RENT':
		rent += int(loanData['loan_amnt'][loanData['member_id'].index(i)])
		numRent+=1
		
	elif ownershipData['home_ownership'][ownershipData['member_id'].index(i)] == 'OWN':
		own += int(loanData['loan_amnt'][loanData['member_id'].index(i)])
		numOwn+=1
		
	else:
		print('Error finding ownership type for id: ' + i)

#average values for loan ammnt
avgMortgage = mortgage/numMortgage
avgRent = rent/numRent
avgOwn = own/numOwn

data = [
	['Ownership Type', 'Average Loan'],
	['Mortage', avgMortgage],
	['Rent', avgRent],
	['Own', avgOwn]
	]

for row in data:
	print(row)


#--- ASK TO INSTALL OPENPYXL ---#

import subprocess
import sys

def install(package):
    subprocess.call([sys.executable, "-m", "pip", "install", package])

#see if openpyxl is installed
try: 
	import openpyxl 
	from openpyxl.chart import BarChart,Reference 
except:
	print()
	print('OPENPYXL does not seem to be Installed... ')
	print('Would You Like to Download OPENPYXL? [Y/N]')
	i = input()
	if i == 'Y' or i == 'y':
		install('openpyxl')
	#try to import again
	try: 
		import openpyxl 
		from openpyxl.chart import BarChart,Reference 
	except:
		print('Unable to install OPENPYXL  :(')
		sys.exit()

#--- PLOT ONTO A GRAPH IN EXCEL UNING OPENPYXL ---#

try:
	  
	# Call a Workbook() function of openpyxl  
	# to create a new blank Workbook object 
	wb = openpyxl.Workbook() 
	  
	# Get workbook active sheet  
	# from the active attribute. 
	sheet = wb.active 
	  
	for row in data: 
		sheet.append(row) 
	  
	  
	# create data for plotting 
	labels = Reference(sheet, min_col = 1, min_row = 2, max_row = len(data)) 
	data = Reference(sheet, min_col = 2, min_row = 1, max_row = len(data)) 
	  
	# Create object of BarChart class 
	chart = BarChart() 
	  
	# adding data to the Bar chart object 
	chart.add_data(data, titles_from_data = True) 

	#add labels to Bar Chart
	chart.set_categories(labels)
	 
	# set the title of the chart 
	chart.title = " OWNERSHIP vs. LOANS "
	  
	# set the title of the x-axis 
	chart.x_axis.title = " OWNERSHIP TYPE "
	  
	# set the title of the y-axis 
	chart.y_axis.title = " AVERAGE LOAN AMOUNT "
	  
	# add chart to the sheet 
	# the top-left corner of a chart 
	# is anchored to cell E2 . 
	sheet.add_chart(chart, "E2") 
	  
	# save the file 
	wb.save("avg_loan.xlsx") 
	
	print()
	print('Excel File has been Created.')

except:
	print('Unable to Plot Graph Using OPENPYXL')
	
try:	
	import os
	os.startfile('avg_loan.xlsx')
except:
	print('Please open avg_loan.xlsx to see Data')